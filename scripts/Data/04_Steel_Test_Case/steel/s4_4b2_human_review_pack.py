"""Create the S4.4b2 human-review pack from compiled workbook outputs.

This module reads only the S4.4b2 workbook snapshot, validation reports and
compiled review CSVs. It does not edit the workbook and does not migrate any
row into executable S4 inputs.
"""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BASE = Path("data/03_Optimisation/inputs/assets/steel/S4/s4_4b2_physical_master_workbook")
COMPILED = BASE / "compiled_review"
OUT = BASE / "human_review_pack"
WORKBOOK = BASE / "steel_c0_c1_physical_model_master.xlsx"
C0 = "C0_current_BF_BOF_reference"
C1 = "C1_phase1_hybrid_BF_BOF_NG_DRP_EAF"


CSV_NAMES = [
    "configurations",
    "carriers",
    "buses",
    "config_components",
    "links",
    "link_ports",
    "operating_constraints",
    "stores",
    "loads",
    "generators",
    "mixing_rules",
    "system_constraints",
    "validation_anchors",
    "evidence_links",
    "modeling_decisions",
    "conflicts_gaps",
]


def _read_csv(name: str) -> list[dict[str, str]]:
    with (COMPILED / f"{name}.csv").open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})


def _norm(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _truthy(value: Any) -> bool:
    return _norm(value).lower() in {"true", "yes", "1"}


def _join_unique(values: list[Any], limit: int | None = None) -> str:
    out: list[str] = []
    for value in values:
        for part in _norm(value).replace("|", ";").split(";"):
            part = part.strip()
            if part and part not in out:
                out.append(part)
    if limit is not None and len(out) > limit:
        return ";".join(out[:limit] + [f"+{len(out) - limit} more"])
    return ";".join(out)


def _compact_ids(ids: list[str], limit: int = 18) -> str:
    ids = [item for item in ids if item]
    if len(ids) <= limit:
        return ";".join(ids)
    return ";".join(ids[:limit] + [f"+{len(ids) - limit} more"])


def _row_status(row: dict[str, str]) -> str:
    keys = [
        "value_status",
        "source_status",
        "executable_status",
        "selected_model_status",
        "source_model_status",
        "selected_hourly_store",
        "profile_type",
    ]
    return _join_unique([row.get(key, "") for key in keys])


def _id_key(sheet_name: str) -> str:
    return {
        "carriers": "carrier_id",
        "buses": "bus_id",
        "config_components": "config_component_id",
        "links": "link_id",
        "link_ports": "port_id",
        "operating_constraints": "constraint_id",
        "stores": "store_id",
        "loads": "load_id",
        "generators": "generator_id",
        "mixing_rules": "mixing_rule_id",
        "system_constraints": "system_constraint_id",
        "validation_anchors": "validation_anchor_id",
        "conflicts_gaps": "gap_or_conflict_id",
        "modeling_decisions": "decision_id",
    }[sheet_name]


def _evidence_index(rows: list[dict[str, str]]) -> dict[tuple[str, str], list[dict[str, str]]]:
    by_entity: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        sheet = _norm(row.get("entity_sheet")).lower()
        entity = _norm(row.get("entity_id"))
        if entity:
            by_entity[(sheet, entity)].append(row)
    return by_entity


def _evidence_refs(
    by_entity: dict[tuple[str, str], list[dict[str, str]]],
    sheet: str,
    entity_ids: str | list[str],
) -> tuple[str, str, str]:
    if isinstance(entity_ids, str):
        ids = [item.strip() for item in entity_ids.replace(",", ";").split(";") if item.strip()]
    else:
        ids = entity_ids
    sources: list[str] = []
    locators: list[str] = []
    strengths: list[str] = []
    for entity_id in ids:
        for row in by_entity.get((sheet.lower(), entity_id), []):
            sources.append(row.get("source_card_id", ""))
            locators.append(row.get("exact_locator", ""))
            strengths.append(row.get("support_strength", ""))
    return _join_unique(sources), _join_unique(locators, limit=6), _join_unique(strengths)


def _load_sources() -> list[dict[str, str]]:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["SOURCES"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, str]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {headers[idx]: values[idx] if idx < len(values) else "" for idx in range(len(headers))}
        if _norm(row.get("source_card_id")):
            rows.append({key: _norm(value) for key, value in row.items()})
    return rows


def _question_rows(
    tables: dict[str, list[dict[str, str]]],
    ev_by_entity: dict[tuple[str, str], list[dict[str, str]]],
) -> list[dict[str, Any]]:
    conflicts = {row["gap_or_conflict_id"]: row for row in tables["conflicts_gaps"]}
    active: dict[str, list[str]] = defaultdict(list)
    inactive: dict[str, list[str]] = defaultdict(list)
    for row in tables["config_components"]:
        target = active if _truthy(row.get("active")) else inactive
        target[row.get("configuration_id", "")].append(row.get("component_id", ""))

    fields = [
        "review_question_id",
        "topic",
        "affected_sheet",
        "affected_entity_ids",
        "current_workbook_status",
        "current_selected_interpretation",
        "alternatives",
        "source_card_ids",
        "exact_locator_if_available",
        "decision_needed",
        "recommended_default",
        "consequence_if_default_selected",
        "consequence_if_not_resolved",
        "blocking_for_s4_4c_execution",
        "blocking_for_s4_4d_economics",
        "human_decision",
        "reviewer_notes",
    ]
    rows: list[dict[str, Any]] = []

    def add(
        qid: str,
        topic: str,
        sheet: str,
        entities: str,
        status: str,
        interpretation: str,
        alternatives: str,
        decision_needed: str,
        default: str,
        default_consequence: str,
        unresolved_consequence: str,
        block_c: str = "yes",
        block_d: str = "yes",
        source_ids: str = "",
        locator: str = "",
    ) -> None:
        if not source_ids or not locator:
            ev_sources, ev_locators, _ = _evidence_refs(ev_by_entity, sheet, entities)
            source_ids = source_ids or ev_sources
            locator = locator or ev_locators
        rows.append(
            {
                "review_question_id": qid,
                "topic": topic,
                "affected_sheet": sheet,
                "affected_entity_ids": entities,
                "current_workbook_status": status,
                "current_selected_interpretation": interpretation,
                "alternatives": alternatives,
                "source_card_ids": source_ids,
                "exact_locator_if_available": locator,
                "decision_needed": decision_needed,
                "recommended_default": default,
                "consequence_if_default_selected": default_consequence,
                "consequence_if_not_resolved": unresolved_consequence,
                "blocking_for_s4_4c_execution": block_c,
                "blocking_for_s4_4d_economics": block_d,
                "human_decision": "",
                "reviewer_notes": "",
            }
        )

    kgf = conflicts.get("CG_C1_KGF_CLOSURE", {})
    add(
        "RQ001",
        "C1 KGF closure governance",
        "CONFLICTS_GAPS",
        "CG_C1_KGF_CLOSURE;coking_plant_1;coking_plant_2;blast_furnace_7",
        "human_review_required; selected default frozen but not migration-approved",
        kgf.get("selected_current_interpretation", "Use frozen default BF7 plus KGF2 closure and retain KGF1."),
        "Close KGF1 instead; close KGF2; keep both coking plants; aggregate coking plants only",
        "Confirm the selected C1 coking-plant closure naming and asset activation.",
        "Retain the frozen default: BF7 and KGF2 closed, KGF1 retained.",
        "C1 topology remains aligned with the current repository freeze and MER hierarchy.",
        "C1 activation and WAG source availability cannot be migrated without a topology decision.",
        source_ids=kgf.get("source_card_ids", ""),
        locator=kgf.get("source_locator", ""),
    )
    add(
        "RQ002",
        "C0 active asset list",
        "CONFIG_COMPONENTS",
        _compact_ids(active[C0]),
        "structural review snapshot; no executable migration approved",
        "C0 includes BF/BOF route, coking, sinter, pellet, downstream, ASU, WAG/steam/Vattenfall interfaces, residual loads and external supplies.",
        "Omit secondary utilities; aggregate interfaces; retain full workbook topology as review-only",
        "Confirm which C0 assets appear in executable S4 inputs and which stay structural/reporting-only.",
        "Use the workbook as physical master and migrate only explicitly approved executable subsets.",
        "Preserves comprehensive topology while avoiding premature executable complexity.",
        "S4.4c migration could miss assets or over-model non-executable utilities.",
    )
    add(
        "RQ003",
        "C1 active and closed asset list",
        "CONFIG_COMPONENTS",
        _compact_ids(active[C1] + inactive[C1]),
        "C1 represented as C0 change set; still human-review gated",
        "C1 adds NG_DRP/EAF/DRI buffer and closes BF7 and KGF2 while retaining the frozen BF/BOF/downstream boundary.",
        "Different retained BF/BOF subset; aggregate old route; DRP/EAF only; separate C1 model",
        "Confirm retained, closed, added and aggregated C1 components before migration.",
        "Use the frozen change-set representation and keep inactive components visible.",
        "Maintains comparability between C0 and C1 and preserves WAG source changes.",
        "C1 dispatch inputs could inherit impossible C0 WAG/material availability.",
    )

    question_specs = [
        ("RQ004", "WAG carriers and buses", "CARRIERS", "BFG;COG;BOFG;process_specific_WAG_blend", "separate carriers/buses present; heating values source-linked where available", "BFG, COG and BOFG remain separate; mixed gas is a process-specific blend output.", "Aggregate WAG into one carrier; model WAG only as validation; keep separate carriers", "Confirm executable migration preserves separate BFG/COG/BOFG balances.", "Preserve separate BFG, COG and BOFG carriers.", "Avoids false fungibility and keeps C1 WAG source changes traceable.", "A single WAG carrier would hide closure effects and sink eligibility."),
        ("RQ005", "WAG mixing stations and sink eligibility", "MIXING_RULES", _compact_ids([r["mixing_rule_id"] for r in tables["mixing_rules"]]), "structural eligibility captured; Wobbe/gas-quality limits missing_not_inferred", "Mixers encode public eligibility for COK1, HSM, PEFA grinding/firing, boilers, STEG11 and Vattenfall.", "Migrate eligibility only; add assumed shares; defer WAG mixing; aggregate sinks", "Approve structural eligibility rows and decide whether missing gas-quality limits block execution.", "Migrate eligibility only after review; keep Wobbe/share limits blocked.", "Allows topology migration without inventing gas-quality physics.", "WAG allocation can become physically invalid or over-flexible."),
        ("RQ006", "WAG holder/storage treatment", "STORES", "COG_gas_holder;BFG_gas_holder;BOFG_gas_holder", "holder rows present; capacities/terminal policies missing or deferred", "Physical holders are recorded separately from selected hourly Store status.", "No hourly Store; finite holders; spill/flare only; ignore holders", "Decide whether any WAG holder is active at hourly resolution and define capacity/terminal policy if yes.", "Do not migrate WAG holders as Stores until capacities and terminal rules are approved.", "Prevents artificial WAG time-shifting.", "Hourly WAG flexibility could be materially overstated."),
        ("RQ007", "Steam bus versus steam store treatment", "STORES", "steam_storage;steam_bus", "steam storage false/deferred; steam is an instantaneous utility bus by default", "Steam is balanced as a utility bus unless explicit steam storage evidence is added.", "Add steam Store; keep bus-only; aggregate steam into residual loads", "Confirm steam remains bus-only unless new evidence supports storage.", "Keep steam as bus-only and block steam storage migration.", "Avoids unsupported intertemporal steam flexibility.", "Steam constraints and boiler dispatch may become non-physical."),
        ("RQ008", "Coke/sinter/pellet store classification", "STORES", "coke_store;sinter_store;pellets_store", "stores recorded; capacities and terminal policies missing_blocker", "Inventory buffers are physically relevant but not executable with unlimited capacity.", "Large finite bounds; sourced finite capacities; no hourly Stores; source-model large stores", "Choose whether each store migrates and define capacity and terminal/anti-arbitrage policy.", "Keep structural do-not-migrate until reviewed.", "Prevents false material arbitrage while retaining the physical record.", "Material buffers could create artificial route flexibility."),
        ("RQ009", "Hot-metal/torpedo buffer treatment", "STORES", "hot_metal_buffer", "candidate hourly store but capacity missing_blocker", "Hot-metal buffering is recorded with terminal neutrality required but no public capacity.", "No store; tight finite buffer; candidate large buffer; route-coupling only", "Decide whether hot metal is modelled as a Store and define capacity/residence policy.", "Do not migrate until capacity or explicit assumption is approved.", "Keeps BF/BOF coupling conservative.", "BF/BOF flexibility could be overstated."),
        ("RQ010", "Slab/WIP/hot-cold state treatment", "STORES", "crude_steel_or_slab_wip_store;hot_slab_transfer_buffer", "slab/WIP candidate value exists; hot/cold state and residence policy unresolved", "Workbook records a candidate slab/WIP buffer and a deferred hot-slab abstraction.", "Single WIP buffer; hot/cold split; residence-time hot charging; no WIP Store", "Approve the slab/WIP representation and whether hot/cold state matters.", "Use a single candidate WIP buffer only after review; defer hot/cold detail unless sourced.", "Allows tractable first migration while preserving thermal-state caveat.", "Downstream flexibility and hot-charging claims could be unsupported."),
        ("RQ011", "Oxygen storage treatment", "STORES", "oxygen_buffer;linde_asu;oxygen_bus", "oxygen buffer recorded but capacity and policy missing_blocker", "ASU produces oxygen; oxygen storage is not executable without capacity evidence.", "No oxygen Store; finite buffer; backup source only; ASU fixed/residual", "Decide whether oxygen storage is active and how to prevent artificial electricity shifting.", "Do not migrate oxygen storage until capacity/terminal policy is approved.", "Prevents unsupported ASU load-shifting flexibility.", "Electricity flexibility may be overstated through oxygen inventory."),
        ("RQ012", "Boiler/steam topology", "LINKS", "boilers_15_16_23_24;boiler_41;STEG11;TG2_steam_turbine", "structural links present; efficiencies, capacities and steam basis unresolved", "Boiler and steam links are topology/interface rows, not executable thermodynamic assets.", "Aggregate boilers; separate boilers; residual steam load; defer steam network", "Confirm executable boiler aggregation and required efficiency/steam-basis assumptions.", "Keep topology only until steam basis and efficiencies are approved.", "Preserves WAG sink topology without unsupported coefficients.", "Steam/WAG balances could become numerically inconsistent."),
        ("RQ013", "Vattenfall/interface topology", "LINKS", "IJ01_interface;VN24_interface;VN25_interface;WAG_VATTENFALL_GENERATORS;flare", "interface rows present; capacity and dispatch responsibility unresolved", "Vattenfall is represented as an interface/sink, not a directly dispatched external plant.", "Dispatch external plant; interface only; validation-only; residual net exchange", "Decide which Vattenfall/interface rows migrate and how capacities/spill are represented.", "Keep as interface topology and block capacity-constrained execution until reviewed.", "Avoids falsely dispatching an external plant while preserving boundary.", "WAG reuse and electricity balance may be misallocated."),
        ("RQ014", "Final-product boundary consistency", "BUSES", "rolled_steel_or_final_product_bus;hot_strip_mill;direct_sheet_plant", "common downstream boundary present; target accumulator is not a flexibility asset", "Both configurations close at the same rolled-steel/final-product boundary.", "End at crude steel; end at slabs; split HSM/DSP products; final-product target", "Confirm the executable final-product balance and downstream treatment.", "Preserve common final-product basis across C0/C1.", "Keeps C0/C1 comparison route-neutral.", "Configuration comparison could mix different product scopes."),
        ("RQ015", "Candidate rows that may migrate", "compiled_review", "candidate/source_supported rows in review pack", "migration candidate CSV generated; approval flags false for Codex", "Only rows without blockers are candidates for later S4 migration.", "Migrate all; migrate none; migrate by sheet/type; require new research first", "Select the candidate subset that can become executable development inputs.", "Approve only a minimal physically coherent subset for the next migration task.", "Creates a controlled migration backlog.", "Migration may be incomplete or methodologically overreaching."),
        ("RQ016", "Rows that must remain non-executable", "compiled_review", "validation anchors;conflicts;missing/redacted/deferred rows", "do-not-migrate CSV generated", "Validation anchors, gaps, conflicts, accounting accumulators and missing blockers are excluded.", "Keep excluded; promote after assumption; reporting only; remove from scope", "Confirm these rows remain non-executable until evidence or assumptions are approved.", "Keep all do-not-migrate rows blocked.", "Maintains provenance discipline.", "Redacted or validation-only values could contaminate executable inputs."),
        ("RQ017", "Missing values requiring research or assumptions", "CONFLICTS_GAPS", _compact_ids([r["gap_or_conflict_id"] for r in tables["conflicts_gaps"]]), "17 gaps/conflicts recorded; several block execution", "Missing capacities, profiles, steam basis, gas-quality limits and terminal policies are explicit blockers.", "Research public sources; define assumptions; keep simplified model; defer subsystem", "Prioritise missing values that block minimal executable C0/C1 physical migration.", "Resolve only the minimum set needed for next stage and keep the rest deferred.", "Focuses review effort on migration-critical gaps.", "The next migration task will lack a defensible input set."),
    ]
    for spec in question_specs:
        add(*spec)
    add(
        "RQ018",
        "Direct PDF inspection provenance gap",
        "SOURCES",
        "Master_Thesis_Report_Athanasiadis*.pdf;Master_Thesis_Mukunda_Badarinath*.pdf",
        "exact PDF filenames not present locally; source-card/register evidence used instead",
        "PDF absence is recorded as a provenance caveat.",
        "Add PDFs and re-extract; accept source-card evidence for review; mark affected rows lower confidence",
        "Decide whether human review can proceed from source-carded locators or requires direct PDF reinspection.",
        "Proceed with human review but block thesis/executable claims on affected values until direct source access is confirmed.",
        "Allows review of structure while preserving provenance caveat.",
        "Source confidence may be insufficient for migration of contested rows.",
        block_c="no",
        block_d="yes",
        source_ids="REPO-ATH-GAS-ADDENDUM;REPO-DEEPSEARCH-MEMOS",
        locator="SOURCES caveat; workbook documentation Source Hierarchy",
    )
    for row in rows:
        for field in fields:
            row.setdefault(field, "")
    return rows


def _do_not_rows(
    tables: dict[str, list[dict[str, str]]],
    ev_by_entity: dict[tuple[str, str], list[dict[str, str]]],
) -> tuple[list[dict[str, Any]], set[tuple[str, str]]]:
    fields = [
        "source_sheet",
        "entity_id",
        "reason_not_to_migrate",
        "current_status",
        "blocking_condition",
        "configuration_scope",
        "source_card_ids",
        "exact_locator_if_available",
        "required_human_action",
        "may_reconsider_after",
        "human_decision",
        "reviewer_notes",
    ]
    rows: list[dict[str, Any]] = []
    keys: set[tuple[str, str]] = set()
    blocked_tokens = {
        "missing_blocker",
        "redacted_or_not_public",
        "validation_only",
        "forbidden",
        "deferred",
        "not_applicable",
        "profile_missing",
        "reporting_only",
    }

    def add(sheet_name: str, entity_id: str, reason: str, row: dict[str, str], blocking: str) -> None:
        key = (sheet_name, entity_id)
        if key in keys:
            return
        keys.add(key)
        source_ids = row.get("source_card_ids", "")
        locator = row.get("source_locator", "")
        if not source_ids or not locator:
            ev_sources, ev_locators, _ = _evidence_refs(ev_by_entity, sheet_name, entity_id)
            source_ids = source_ids or ev_sources
            locator = locator or ev_locators
        rows.append(
            {
                "source_sheet": sheet_name,
                "entity_id": entity_id,
                "reason_not_to_migrate": reason,
                "current_status": _row_status(row),
                "blocking_condition": blocking,
                "configuration_scope": row.get("configuration_scope")
                or row.get("configuration_id")
                or row.get("configuration_relevance", ""),
                "source_card_ids": source_ids,
                "exact_locator_if_available": locator,
                "required_human_action": "Keep non-executable until reviewed.",
                "may_reconsider_after": "source evidence or explicit human assumption approved",
                "human_decision": "",
                "reviewer_notes": "",
            }
        )

    data_sheets = [
        "carriers",
        "buses",
        "links",
        "link_ports",
        "operating_constraints",
        "stores",
        "loads",
        "generators",
        "mixing_rules",
        "system_constraints",
    ]
    for sheet_name in data_sheets:
        id_key = _id_key(sheet_name)
        for row in tables[sheet_name]:
            entity_id = row.get(id_key, "")
            status = _row_status(row).lower()
            if any(token in status for token in blocked_tokens):
                add(sheet_name, entity_id, "Workbook status blocks executable migration.", row, status)
            if sheet_name == "stores":
                store_class = row.get("store_class", "").lower()
                if row.get("terminal_inventory_rule", "").lower() == "missing":
                    add(sheet_name, entity_id, "Store terminal policy is missing.", row, "missing terminal inventory policy")
                if store_class in {"accounting accumulator", "target accumulator", "technical source-model workaround"} or row.get("selected_hourly_store", "").lower() == "false":
                    add(sheet_name, entity_id, "Store row is not an executable physical flexibility buffer.", row, store_class or "selected_hourly_store=false")
            if sheet_name == "loads" and row.get("profile_type", "").lower() == "profile_missing":
                add(sheet_name, entity_id, "Load profile/allocation is missing.", row, "profile_missing")
            if sheet_name == "mixing_rules" and row.get("eligible", "").lower() == "false":
                add(sheet_name, entity_id, "Ineligible carrier/mixer pair must not become an executable eligibility row.", row, "eligible=false")

    for row in tables["validation_anchors"]:
        add("validation_anchors", row.get("validation_anchor_id", ""), "Validation anchor is not an executable input.", row, "validation_only")
    for row in tables["conflicts_gaps"]:
        add("conflicts_gaps", row.get("gap_or_conflict_id", ""), "Unresolved gap/conflict must be resolved or explicitly accepted before migration.", row, "human_review_required")
    for row in tables["modeling_decisions"]:
        add("modeling_decisions", row.get("decision_id", ""), "Modeling decision is governance metadata, not an executable input row.", row, "decision_metadata")

    for row in rows:
        for field in fields:
            row.setdefault(field, "")
    return rows, keys


def _migration_rows(
    tables: dict[str, list[dict[str, str]]],
    ev_by_entity: dict[tuple[str, str], list[dict[str, str]]],
    blocked_keys: set[tuple[str, str]],
) -> list[dict[str, Any]]:
    target_table = {
        "carriers": "future_carrier_properties.csv",
        "buses": "future_balance_nodes.csv",
        "config_components": "configuration_assets.csv",
        "links": "future_physical_links.csv",
        "link_ports": "future_process_coefficients_or_ports.csv",
        "operating_constraints": "future_asset_operating_constraints.csv",
        "stores": "buffers_and_stores.csv",
        "generators": "future_external_supplies.csv",
        "mixing_rules": "future_wag_mixing_rules.csv",
        "system_constraints": "future_system_constraints.csv",
    }
    target_fields = {
        "carriers": "carrier_id;balance_unit;heating_value_fields;storable_flags;active_C0;active_C1",
        "buses": "bus_id;carrier_id;network_layer;balance_unit;selected_model_role",
        "config_components": "configuration_id;component_type;component_id;active;retained_added_closed;controllable",
        "links": "link_id;process_family;throughput_basis;controllability;representation_status",
        "link_ports": "link_id;bus_id;direction;coefficient_fields_if_populated;coefficient_basis",
        "operating_constraints": "component_id;constraint_family;lower_value;central_value;upper_value;unit;basis",
        "stores": "store_id;bus_id;capacity_fields;initial_policy;terminal_policy;charge_discharge_fields",
        "generators": "generator_id;bus_id;available;capacity_bound;availability_profile",
        "mixing_rules": "mixer_link_id;output_bus_id;input_carrier_id;eligible;share_fields;gas_quality_fields_if_sourced",
        "system_constraints": "system_constraint_id;constraint_family;scope;selected_for_execution_after_approval",
    }
    rows: list[dict[str, Any]] = []
    for sheet_name in target_table:
        id_key = _id_key(sheet_name)
        for row in tables[sheet_name]:
            entity_id = row.get(id_key, "")
            if (sheet_name, entity_id) in blocked_keys:
                continue
            status = _row_status(row).lower()
            candidate = any(token in status for token in ["candidate", "source_supported", "development_review", "assumption"])
            if sheet_name == "config_components":
                candidate = bool(row.get("source_model_status") or row.get("selected_model_status"))
                if row.get("topology_decision_id") == "CG_C1_KGF_CLOSURE":
                    candidate = False
            if sheet_name == "link_ports":
                candidate = row.get("source_status", "").lower() in {"candidate", "source_supported", "assumption"}
            if sheet_name == "mixing_rules":
                candidate = row.get("eligible", "").lower() == "true" and row.get("source_status", "").lower() == "source_supported"
            if not candidate:
                continue
            _, _, strength = _evidence_refs(ev_by_entity, sheet_name, entity_id)
            missing: list[str] = []
            if sheet_name == "mixing_rules" and "missing" in row.get("Wobbe_constraint_status", "").lower():
                missing.append("Wobbe/gas-quality quantitative limits")
            if sheet_name == "stores":
                if not any(row.get(key) for key in ["capacity_lower", "capacity_central", "capacity_upper"]):
                    missing.append("capacity")
                if not row.get("terminal_inventory_rule"):
                    missing.append("terminal_inventory_rule")
            if sheet_name == "link_ports" and not any(row.get(key) for key in ["coefficient_lower", "coefficient_central", "coefficient_upper"]):
                missing.append("numeric coefficient if required by executable formulation")
            reason = "Source-supported/candidate row can be considered for controlled S4 migration after human approval."
            if sheet_name == "operating_constraints":
                reason = "Candidate numeric operating constraint has explicit provenance and can be reviewed for executable use."
            elif sheet_name == "stores":
                reason = "Candidate buffer has explicit capacity/policy fields and can be reviewed for controlled execution."
            elif sheet_name == "mixing_rules":
                reason = "Public structural WAG eligibility can be migrated separately from missing quantitative gas-quality limits."
            rows.append(
                {
                    "source_sheet": sheet_name,
                    "entity_id": entity_id,
                    "target_s4_table": target_table[sheet_name],
                    "target_fields": target_fields[sheet_name],
                    "configuration_scope": row.get("configuration_scope") or row.get("configuration_id") or f"{row.get('active_C0', '')}/{row.get('active_C1', '')}",
                    "current_status": _row_status(row),
                    "proposed_status_after_approval": "approved_development_input_candidate",
                    "reason_for_candidate_migration": reason,
                    "evidence_strength": strength or row.get("source_status") or row.get("value_status") or "review_required",
                    "missing_fields": ";".join(missing),
                    "human_review_required": "true",
                    "codex_may_decide": "false",
                }
            )
    return rows


def _wag_rows(tables: dict[str, list[dict[str, str]]]) -> list[dict[str, Any]]:
    carriers = {row["carrier_id"]: row for row in tables["carriers"]}
    buses = {row["bus_id"]: row for row in tables["buses"]}
    stores = {row["store_id"]: row for row in tables["stores"]}
    carrier_to_buses: dict[str, list[str]] = defaultdict(list)
    ports_by_bus: dict[str, list[dict[str, str]]] = defaultdict(list)
    for bus in tables["buses"]:
        carrier_to_buses[bus.get("carrier_id", "")].append(bus.get("bus_id", ""))
    for port in tables["link_ports"]:
        ports_by_bus[port.get("bus_id", "")].append(port)

    rows: list[dict[str, Any]] = []
    for carrier_id in ["BFG", "COG", "BOFG", "process_specific_WAG_blend"]:
        carrier = carriers.get(carrier_id, {})
        producers: list[str] = []
        flare: list[str] = []
        for bus_id in carrier_to_buses.get(carrier_id, []):
            for port in ports_by_bus[bus_id]:
                link_id = port.get("link_id", "")
                if port.get("direction") == "output":
                    producers.append(link_id)
                if "flare" in link_id.lower():
                    flare.append(link_id)
        sinks = [row for row in tables["mixing_rules"] if row.get("input_carrier_id") == carrier_id and row.get("eligible") == "true"]
        holder_ids = [
            store["store_id"]
            for store in tables["stores"]
            if buses.get(store.get("bus_id", ""), {}).get("carrier_id") == carrier_id
        ]
        rows.append(
            {
                "record_type": "carrier",
                "carrier_or_mixer_id": carrier_id,
                "BFG_COG_BOFG_or_mixed_gas_carrier": carrier_id,
                "generation_source": _join_unique(producers),
                "active_C0": carrier.get("active_C0", ""),
                "active_C1": carrier.get("active_C1", ""),
                "sink_eligibility": _join_unique([f"{row.get('mixer_link_id')}->{row.get('output_bus_id')}" for row in sinks], limit=12),
                "storage_holder_status": _join_unique([f"{sid}:{stores.get(sid, {}).get('selected_hourly_store', '')}/{stores.get(sid, {}).get('value_status', '')}" for sid in holder_ids]) or "none recorded",
                "flare_spill_status": _join_unique(flare) or ("flare link present for WAG network" if carrier_id in {"BFG", "COG", "BOFG"} else "not a primary flare carrier"),
                "LHV_basis": _join_unique([carrier.get("LHV_or_HHV_basis"), carrier.get("heating_value_lower"), carrier.get("heating_value_central"), carrier.get("heating_value_upper"), carrier.get("heating_value_unit")]),
                "missing_Wobbe_gas_quality_limits": "yes",
                "migration_recommendation": "structural carrier may migrate after review; holder storage and Wobbe limits remain blocked unless approved",
                "review_decision_needed": "Confirm separate carrier balance, C1 source availability, holder treatment and gas-quality handling.",
            }
        )
    for mixer_id in sorted({row.get("mixer_link_id", "") for row in tables["mixing_rules"] if row.get("mixer_link_id")}):
        rules = [row for row in tables["mixing_rules"] if row.get("mixer_link_id") == mixer_id]
        eligible = [row for row in rules if row.get("eligible") == "true"]
        blocked = [row for row in rules if row.get("eligible") == "false"]
        rows.append(
            {
                "record_type": "mixer",
                "carrier_or_mixer_id": mixer_id,
                "BFG_COG_BOFG_or_mixed_gas_carrier": _join_unique([row.get("input_carrier_id") for row in rules]),
                "generation_source": "mixer/controller, not a gas producer",
                "active_C0": "review_required",
                "active_C1": "review_required",
                "sink_eligibility": f"eligible={_join_unique([row.get('input_carrier_id') for row in eligible])}; forbidden={_join_unique([row.get('input_carrier_id') for row in blocked])}",
                "storage_holder_status": "not a Store; controller row",
                "flare_spill_status": "use flare/spill closure only where connected downstream",
                "LHV_basis": _join_unique([row.get("energy_equivalence_basis") for row in rules]),
                "missing_Wobbe_gas_quality_limits": "yes" if any("missing" in row.get("Wobbe_constraint_status", "").lower() for row in rules) else "no",
                "migration_recommendation": "structural eligibility may migrate after review; quantitative gas-quality limits must not be invented",
                "review_decision_needed": "Approve eligible input carriers and decide whether absent share/Wobbe limits block execution.",
            }
        )
    return rows


def _store_rows(tables: dict[str, list[dict[str, str]]], migration_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    buses = {row["bus_id"]: row for row in tables["buses"]}
    component_active: dict[tuple[str, str], dict[str, str]] = defaultdict(dict)
    for row in tables["config_components"]:
        component_active[(row.get("component_type", ""), row.get("component_id", ""))][row.get("configuration_id", "")] = row.get("active", "")
    migration_store_ids = {row["entity_id"] for row in migration_rows if row["source_sheet"] == "stores"}
    rows: list[dict[str, Any]] = []
    for store in tables["stores"]:
        capacity_known = "yes" if any(store.get(key) for key in ["capacity_lower", "capacity_central", "capacity_upper"]) else "no"
        initial_known = "yes" if store.get("initial_inventory_rule") and store.get("initial_inventory_rule") != "missing" else "no"
        terminal_known = "yes" if store.get("terminal_inventory_rule") and store.get("terminal_inventory_rule") != "missing" else "no"
        risk = "high" if store.get("selected_hourly_store") in {"candidate", "deferred"} and (capacity_known == "no" or terminal_known == "no") else "low"
        if store.get("store_class") in {"accounting accumulator", "target accumulator", "technical source-model workaround"}:
            risk = "high_if_treated_as_physical_flexibility"
        rows.append(
            {
                "store_id": store["store_id"],
                "carrier": buses.get(store.get("bus_id", ""), {}).get("carrier_id", ""),
                "physical_asset_exists": store.get("physical_asset_exists", ""),
                "source_model_store": store.get("source_model_store", ""),
                "selected_hourly_store": store.get("selected_hourly_store", ""),
                "active_C0": component_active.get(("Store", store["store_id"]), {}).get(C0, ""),
                "active_C1": component_active.get(("Store", store["store_id"]), {}).get(C1, ""),
                "capacity_known": capacity_known,
                "initial_policy_known": initial_known,
                "terminal_policy_known": terminal_known,
                "risk_of_fake_flexibility": risk,
                "migration_recommendation": "candidate_after_human_review" if store["store_id"] in migration_store_ids else "do_not_migrate_until_review",
                "review_decision_needed": "Approve physical hourly Store status, capacity, initial rule and terminal rule before migration.",
            }
        )
    return rows


def _topology_rows(tables: dict[str, list[dict[str, str]]]) -> list[dict[str, Any]]:
    rows = tables["config_components"]
    links_c0_active = [row["component_id"] for row in rows if row.get("configuration_id") == C0 and row.get("component_type") == "Link" and row.get("active") == "true"]
    links_c1_inactive = [row["component_id"] for row in rows if row.get("configuration_id") == C1 and row.get("component_type") == "Link" and row.get("active") == "false"]
    retained = [row["component_id"] for row in rows if row.get("configuration_id") == C1 and row.get("component_type") == "Link" and row.get("retained_added_closed") == "retained"]
    added = [row["component_id"] for row in rows if row.get("configuration_id") == C1 and row.get("component_type") == "Link" and row.get("retained_added_closed") == "added"]
    closed = [row["component_id"] for row in rows if row.get("configuration_id") == C1 and row.get("component_type") == "Link" and row.get("retained_added_closed") == "closed"]
    hierarchy = "MER/formal Tata public evidence > Athanasiadis modelling architecture > Badarinath secondary precedent > repository freezes > deepsearch discovery memos"
    return [
        {
            "review_topic": "active C0 assets",
            "configuration_scope": C0,
            "entity_group": "active Link components",
            "entity_ids": _compact_ids(links_c0_active),
            "current_workbook_status": "structural master; not executable",
            "selected_interpretation": "C0 current BF/BOF reference route plus utilities, WAG/steam interfaces and downstream boundary.",
            "source_hierarchy_used": hierarchy,
            "unresolved_conflicts": "",
            "decision_needed": "Confirm migration subset for executable C0 dev inputs.",
            "recommended_default": "Migrate only approved executable core route plus required utilities.",
            "blocking_for_migration": "yes",
            "reviewer_notes": "",
        },
        {
            "review_topic": "closed/inactive C1 assets",
            "configuration_scope": C1,
            "entity_group": "closed Link components",
            "entity_ids": _compact_ids(closed + links_c1_inactive),
            "current_workbook_status": "frozen default visible; KGF closure requires human review",
            "selected_interpretation": "BF7 and coking_plant_2/KGF2 closed in default C1.",
            "source_hierarchy_used": hierarchy,
            "unresolved_conflicts": "CG_C1_KGF_CLOSURE",
            "decision_needed": "Confirm closed assets and WAG generation consequences.",
            "recommended_default": "Retain frozen default unless human review changes the topology.",
            "blocking_for_migration": "yes",
            "reviewer_notes": "",
        },
        {
            "review_topic": "retained C1 BF-BOF assets",
            "configuration_scope": C1,
            "entity_group": "retained Link components",
            "entity_ids": _compact_ids(retained),
            "current_workbook_status": "retained route represented as change set over C0",
            "selected_interpretation": "C1 remains hybrid and retains selected BF/BOF/downstream/utility assets.",
            "source_hierarchy_used": hierarchy,
            "unresolved_conflicts": "CG_C1_KGF_CLOSURE affects retained coking source naming",
            "decision_needed": "Confirm retained-route coupling and executable utilities.",
            "recommended_default": "Keep retained route visible but migrate only required executable subset.",
            "blocking_for_migration": "yes",
            "reviewer_notes": "",
        },
        {
            "review_topic": "added C1 DRP/EAF assets",
            "configuration_scope": C1,
            "entity_group": "added Link components",
            "entity_ids": _compact_ids(added),
            "current_workbook_status": "candidate numeric rows present for NG_DRP and EAF",
            "selected_interpretation": "C1 adds NG_DRP, DRI buffer and EAF integrated into common downstream boundary.",
            "source_hierarchy_used": hierarchy,
            "unresolved_conflicts": "candidate coefficients require approval; not Tata validated",
            "decision_needed": "Approve candidate DRP/EAF capacities, coefficients and operating constraints or replace with assumptions.",
            "recommended_default": "Use candidate rows only for development after explicit approval.",
            "blocking_for_migration": "yes",
            "reviewer_notes": "",
        },
        {
            "review_topic": "downstream/final-product boundary",
            "configuration_scope": "C0 and C1",
            "entity_group": "boundary buses and downstream links",
            "entity_ids": "continuous_caster_or_slab_conversion;hot_strip_mill;direct_sheet_plant;rolled_steel_or_final_product_bus",
            "current_workbook_status": "common boundary represented; detailed downstream capacities unresolved",
            "selected_interpretation": "Both configurations use the same final-product boundary for comparability.",
            "source_hierarchy_used": hierarchy,
            "unresolved_conflicts": "CG_DOWNSTREAM_CAPACITIES",
            "decision_needed": "Confirm whether executable migration ends at liquid steel, slab or final product.",
            "recommended_default": "Keep final-product boundary and defer detailed downstream constraints if not sourced.",
            "blocking_for_migration": "yes",
            "reviewer_notes": "",
        },
        {
            "review_topic": "source hierarchy used",
            "configuration_scope": "C0 and C1",
            "entity_group": "all topology decisions",
            "entity_ids": "all",
            "current_workbook_status": "source hierarchy documented; direct PDFs absent",
            "selected_interpretation": hierarchy,
            "source_hierarchy_used": hierarchy,
            "unresolved_conflicts": _join_unique([row["gap_or_conflict_id"] for row in tables["conflicts_gaps"] if row.get("human_review_required") == "true"], limit=8),
            "decision_needed": "Confirm this hierarchy remains governing for migration.",
            "recommended_default": "Keep repository freeze hierarchy until source documents are rechecked.",
            "blocking_for_migration": "yes",
            "reviewer_notes": "",
        },
    ]


def _provenance_rows(sources: list[dict[str, str]], tables: dict[str, list[dict[str, str]]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    def add(pid: str, gap_type: str, affected: str, sheet: str, description: str, source_ids: str, locator: str, status: str, action: str, block: str) -> None:
        rows.append(
            {
                "provenance_gap_id": pid,
                "gap_type": gap_type,
                "affected_source_or_entity": affected,
                "affected_sheet": sheet,
                "description": description,
                "source_card_ids": source_ids,
                "exact_locator_if_available": locator,
                "current_status": status,
                "human_action_needed": action,
                "blocking_for_migration": block,
            }
        )

    add("PG001", "source_pdf_not_directly_inspected", "Master_Thesis_Report_Athanasiadis*.pdf", "SOURCES", "Requested Athanasiadis PDF filename was not present in the repository snapshot used for workbook construction.", "REPO-ATH-GAS-ADDENDUM;REPO-DEEPSEARCH-MEMOS", "SOURCES caveat; workbook documentation Source Hierarchy", "source-card/addendum evidence used", "Provide PDF or confirm source-card evidence is sufficient for review.", "yes_for_affected_rows")
    add("PG002", "source_pdf_not_directly_inspected", "Master_Thesis_Mukunda_Badarinath*.pdf", "SOURCES", "Requested Badarinath PDF filename was not present in the repository snapshot used for workbook construction.", "REPO-DEEPSEARCH-MEMOS", "SOURCES caveat; workbook documentation Source Hierarchy", "registered evidence and policy summaries used", "Provide PDF or confirm registered evidence is sufficient for review.", "yes_for_affected_rows")
    idx = 3
    for source in sources:
        caveat = source.get("caveat", "").lower()
        status = source.get("source_review_status", "").lower()
        if "pdf was not present" in caveat or "source-card" in caveat or "addendum" in source.get("source_title", "").lower() or status in {"locator_incomplete", "read_for_task"}:
            add(f"PG{idx:03d}", "source_card_or_register_evidence_only", source["source_card_id"], "SOURCES", "Source row is based on repository source card/register/deepsearch evidence or a task-level read, not necessarily direct primary PDF reinspection.", source["source_card_id"], source.get("exact_locator", ""), source.get("source_review_status", ""), "Review whether this evidence level is sufficient for migration-critical rows.", "review_required")
            idx += 1
        if not source.get("source_url_or_doi"):
            add(f"PG{idx:03d}", "missing_public_url_or_doi", source["source_card_id"], "SOURCES", "No public URL/DOI is recorded. This may be acceptable for repo-local governance files but should be reviewed for external source claims.", source["source_card_id"], source.get("exact_locator", ""), source.get("source_review_status", ""), "Add URL/DOI where public and stable; otherwise keep repo-relative identifier.", "no_for_repo_governance_yes_for_external_claims")
            idx += 1
    for row in tables["conflicts_gaps"]:
        if row.get("human_review_required") == "true":
            add(f"PG{idx:03d}", "unresolved_conflict_or_gap", row.get("gap_or_conflict_id", ""), "CONFLICTS_GAPS", row.get("description", ""), row.get("source_card_ids", ""), row.get("source_locator", ""), _row_status(row) or "human_review_required", "Resolve, explicitly accept, or keep blocked before migration.", "yes")
            idx += 1
    return rows


def build_review_pack() -> dict[str, Any]:
    OUT.mkdir(parents=True, exist_ok=True)
    tables = {name: _read_csv(name) for name in CSV_NAMES}
    sources = _load_sources()
    ev_by_entity = _evidence_index(tables["evidence_links"])

    question_rows = _question_rows(tables, ev_by_entity)
    do_not_rows, blocked_keys = _do_not_rows(tables, ev_by_entity)
    migration_rows = _migration_rows(tables, ev_by_entity, blocked_keys)
    wag_rows = _wag_rows(tables)
    store_rows = _store_rows(tables, migration_rows)
    topology_rows = _topology_rows(tables)
    provenance_rows = _provenance_rows(sources, tables)

    stage_gate = {
        "stage": "s4_4b2_human_review_pack",
        "decision": "ready_for_human_review",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "workbook_path": str(WORKBOOK).replace("\\", "/"),
        "review_pack_dir": str(OUT).replace("\\", "/"),
        "forbidden_decisions": [
            "ready_for_model_migration",
            "ready_for_unified_model_execution",
            "thesis_usable",
            "Tata-validated",
        ],
        "counts": {
            "review_questions": len(question_rows),
            "migration_candidate_rows": len(migration_rows),
            "do_not_migrate_rows": len(do_not_rows),
            "wag_review_rows": len(wag_rows),
            "store_buffer_review_rows": len(store_rows),
            "topology_review_rows": len(topology_rows),
            "source_provenance_gap_rows": len(provenance_rows),
        },
        "blockers_before_migration": [
            "human review of KGF1/KGF2 closure governance",
            "approval of executable C0/C1 asset subset",
            "explicit WAG mixing/holder/gas-quality decisions",
            "store and terminal inventory policies",
            "steam and oxygen storage treatment",
            "source provenance caveat for absent local thesis PDFs",
        ],
        "migration_allowed": False,
        "codex_may_decide_migration": False,
    }

    _write_csv(OUT / "s4_4b2_human_review_questions.csv", [
        "review_question_id",
        "topic",
        "affected_sheet",
        "affected_entity_ids",
        "current_workbook_status",
        "current_selected_interpretation",
        "alternatives",
        "source_card_ids",
        "exact_locator_if_available",
        "decision_needed",
        "recommended_default",
        "consequence_if_default_selected",
        "consequence_if_not_resolved",
        "blocking_for_s4_4c_execution",
        "blocking_for_s4_4d_economics",
        "human_decision",
        "reviewer_notes",
    ], question_rows)
    _write_csv(OUT / "s4_4b2_migration_candidate_rows.csv", [
        "source_sheet",
        "entity_id",
        "target_s4_table",
        "target_fields",
        "configuration_scope",
        "current_status",
        "proposed_status_after_approval",
        "reason_for_candidate_migration",
        "evidence_strength",
        "missing_fields",
        "human_review_required",
        "codex_may_decide",
    ], migration_rows)
    _write_csv(OUT / "s4_4b2_do_not_migrate_rows.csv", [
        "source_sheet",
        "entity_id",
        "reason_not_to_migrate",
        "current_status",
        "blocking_condition",
        "configuration_scope",
        "source_card_ids",
        "exact_locator_if_available",
        "required_human_action",
        "may_reconsider_after",
        "human_decision",
        "reviewer_notes",
    ], do_not_rows)
    _write_csv(OUT / "s4_4b2_wag_review_pack.csv", [
        "record_type",
        "carrier_or_mixer_id",
        "BFG_COG_BOFG_or_mixed_gas_carrier",
        "generation_source",
        "active_C0",
        "active_C1",
        "sink_eligibility",
        "storage_holder_status",
        "flare_spill_status",
        "LHV_basis",
        "missing_Wobbe_gas_quality_limits",
        "migration_recommendation",
        "review_decision_needed",
    ], wag_rows)
    _write_csv(OUT / "s4_4b2_store_buffer_review_pack.csv", [
        "store_id",
        "carrier",
        "physical_asset_exists",
        "source_model_store",
        "selected_hourly_store",
        "active_C0",
        "active_C1",
        "capacity_known",
        "initial_policy_known",
        "terminal_policy_known",
        "risk_of_fake_flexibility",
        "migration_recommendation",
        "review_decision_needed",
    ], store_rows)
    _write_csv(OUT / "s4_4b2_c0_c1_topology_review_pack.csv", [
        "review_topic",
        "configuration_scope",
        "entity_group",
        "entity_ids",
        "current_workbook_status",
        "selected_interpretation",
        "source_hierarchy_used",
        "unresolved_conflicts",
        "decision_needed",
        "recommended_default",
        "blocking_for_migration",
        "reviewer_notes",
    ], topology_rows)
    _write_csv(OUT / "s4_4b2_source_provenance_gaps.csv", [
        "provenance_gap_id",
        "gap_type",
        "affected_source_or_entity",
        "affected_sheet",
        "description",
        "source_card_ids",
        "exact_locator_if_available",
        "current_status",
        "human_action_needed",
        "blocking_for_migration",
    ], provenance_rows)
    with (OUT / "s4_4b2_review_stage_gate.json").open("w", encoding="utf-8") as handle:
        json.dump(stage_gate, handle, indent=2)
    _write_csv(OUT / "s4_4b2_review_stage_gate.csv", [
        "stage",
        "decision",
        "migration_allowed",
        "codex_may_decide_migration",
        "review_questions",
        "migration_candidate_rows",
        "do_not_migrate_rows",
        "source_provenance_gap_rows",
    ], [{
        "stage": stage_gate["stage"],
        "decision": stage_gate["decision"],
        "migration_allowed": "false",
        "codex_may_decide_migration": "false",
        "review_questions": len(question_rows),
        "migration_candidate_rows": len(migration_rows),
        "do_not_migrate_rows": len(do_not_rows),
        "source_provenance_gap_rows": len(provenance_rows),
    }])

    _validate_pack(question_rows, migration_rows, do_not_rows)
    return stage_gate


def _validate_pack(
    question_rows: list[dict[str, Any]],
    migration_rows: list[dict[str, Any]],
    do_not_rows: list[dict[str, Any]],
) -> None:
    question_ids = [row["review_question_id"] for row in question_rows]
    if len(question_ids) != len(set(question_ids)):
        raise ValueError("duplicate review_question_id detected")
    bad_flags = [
        row for row in migration_rows
        if row["human_review_required"] != "true" or row["codex_may_decide"] != "false"
    ]
    if bad_flags:
        raise ValueError("migration candidate approval flags are invalid")
    candidate_keys = {(row["source_sheet"], row["entity_id"]) for row in migration_rows}
    blocked_keys = {(row["source_sheet"], row["entity_id"]) for row in do_not_rows}
    overlap = candidate_keys & blocked_keys
    if overlap:
        raise ValueError(f"do-not/migration overlap detected: {sorted(overlap)[:5]}")


def main() -> None:
    stage_gate = build_review_pack()
    print(json.dumps({"decision": stage_gate["decision"], "counts": stage_gate["counts"]}, indent=2))


if __name__ == "__main__":
    main()
